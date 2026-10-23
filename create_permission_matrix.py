import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def create_permission_matrix_excel(output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bảng phân quyền 5 vai trò"

    # Font styles
    font_title = Font(name='Arial', size=13, bold=True, color='1E293B')
    font_subtitle = Font(name='Arial', size=9, italic=True, color='64748B')
    font_group = Font(name='Arial', size=10, bold=True, color='1E3A8A')
    font_header_role = Font(name='Arial', size=10, bold=True, color='FFFFFF')
    font_header_action = Font(name='Arial', size=9, bold=True, color='FFFFFF')
    font_true = Font(name='Arial', size=9, bold=True, color='047857')
    font_false = Font(name='Arial', size=9, color='94A3B8')
    font_text = Font(name='Arial', size=9, color='1E293B')
    font_note = Font(name='Arial', size=8, italic=True, color='475569')

    # Fills
    fill_admin = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')     # Deep Navy
    fill_so = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')        # Teal
    fill_truong = PatternFill(start_color='D97706', end_color='D97706', fill_type='solid')    # Amber
    fill_dn = PatternFill(start_color='7C3AED', end_color='7C3AED', fill_type='solid')        # Purple
    fill_khach = PatternFill(start_color='475569', end_color='475569', fill_type='solid')     # Slate
    fill_main = PatternFill(start_color='0F172A', end_color='0F172A', fill_type='solid')      # Dark Slate

    fill_action_admin = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
    fill_action_so = PatternFill(start_color='14B8A6', end_color='14B8A6', fill_type='solid')
    fill_action_truong = PatternFill(start_color='F59E0B', end_color='F59E0B', fill_type='solid')
    fill_action_dn = PatternFill(start_color='8B5CF6', end_color='8B5CF6', fill_type='solid')
    fill_action_khach = PatternFill(start_color='64748B', end_color='64748B', fill_type='solid')

    fill_group_row = PatternFill(start_color='F1F5F9', end_color='F1F5F9', fill_type='solid')
    fill_zebra = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
    fill_true = PatternFill(start_color='ECFDF5', end_color='ECFDF5', fill_type='solid')

    thin_border_side = Side(border_style='thin', color='CBD5E1')
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_left = Alignment(horizontal='left', vertical='center', wrap_text=True)

    # Title
    ws.merge_cells('A1:X1')
    ws['A1'] = "MA TRẬN PHÂN QUYỀN CHỨC NĂNG DỰ ÁN BẢN ĐỒ SỐ GDNN & VIỆC LÀM NINH BÌNH"
    ws['A1'].font = font_title
    ws['A1'].alignment = align_left
    ws.row_dimensions[1].height = 26

    ws.merge_cells('A2:X2')
    ws['A2'] = "5 Vai trò: 1. Admin (Quản trị hệ thống) | 2. Sở Giáo dục (Cơ quan QLNN) | 3. Nhà trường (Cơ sở GDNN) | 4. Doanh nghiệp (Tuyển dụng) | 5. Khách (Người dân / Học sinh)"
    ws['A2'].font = font_subtitle
    ws['A2'].alignment = align_left
    ws.row_dimensions[2].height = 18

    # Header Row 1 (Role level)
    # Col 1: STT, Col 2: Danh sách chức năng
    ws.merge_cells('A4:A5')
    ws['A4'] = "STT"
    ws['A4'].fill = fill_main
    ws['A4'].font = font_header_role
    ws['A4'].alignment = align_center
    ws['A4'].border = border_cell

    ws.merge_cells('B4:B5')
    ws['B4'] = "Danh sách chức năng hệ thống"
    ws['B4'].fill = fill_main
    ws['B4'].font = font_header_role
    ws['B4'].alignment = align_center
    ws['B4'].border = border_cell

    # Role 1: Admin (Cols C-G: Xem, Thêm, Sửa, Xóa, Phê duyệt)
    ws.merge_cells('C4:G4')
    ws['C4'] = "Admin (Quản trị viên)"
    ws['C4'].fill = fill_admin
    ws['C4'].font = font_header_role
    ws['C4'].alignment = align_center
    ws['C4'].border = border_cell

    # Role 2: Sở Giáo dục (Cols H-L: Xem, Thêm, Sửa, Xóa, Phê duyệt)
    ws.merge_cells('H4:L4')
    ws['H4'] = "Sở Giáo dục / Quản lý"
    ws['H4'].fill = fill_so
    ws['H4'].font = font_header_role
    ws['H4'].alignment = align_center
    ws['H4'].border = border_cell

    # Role 3: Nhà trường (Cols M-P: Xem, Thêm, Sửa, Xóa)
    ws.merge_cells('M4:P4')
    ws['M4'] = "Nhà trường (Cơ sở GDNN)"
    ws['M4'].fill = fill_truong
    ws['M4'].font = font_header_role
    ws['M4'].alignment = align_center
    ws['M4'].border = border_cell

    # Role 4: Doanh nghiệp (Cols Q-T: Xem, Thêm, Sửa, Xóa)
    ws.merge_cells('Q4:T4')
    ws['Q4'] = "Doanh nghiệp (Tuyển dụng)"
    ws['Q4'].fill = fill_dn
    ws['Q4'].font = font_header_role
    ws['Q4'].alignment = align_center
    ws['Q4'].border = border_cell

    # Role 5: Khách (Cols U-W: Xem, Thêm, Sửa)
    ws.merge_cells('U4:W4')
    ws['U4'] = "Khách (Người dân / Học sinh)"
    ws['U4'].fill = fill_khach
    ws['U4'].font = font_header_role
    ws['U4'].alignment = align_center
    ws['U4'].border = border_cell

    # Col X: Ghi chú
    ws.merge_cells('X4:X5')
    ws['X4'] = "Phạm vi dữ liệu & Ghi chú nghiệp vụ"
    ws['X4'].fill = fill_main
    ws['X4'].font = font_header_role
    ws['X4'].alignment = align_center
    ws['X4'].border = border_cell

    ws.row_dimensions[4].height = 24

    # Header Row 2 (Action level)
    actions = [
        # Admin (C-G)
        ('Xem', fill_action_admin), ('Thêm', fill_action_admin), ('Sửa', fill_action_admin), ('Xóa', fill_action_admin), ('Duyệt', fill_action_admin),
        # Sở (H-L)
        ('Xem', fill_action_so), ('Thêm', fill_action_so), ('Sửa', fill_action_so), ('Xóa', fill_action_so), ('Duyệt', fill_action_so),
        # Trường (M-P)
        ('Xem', fill_action_truong), ('Thêm', fill_action_truong), ('Sửa', fill_action_truong), ('Xóa', fill_action_truong),
        # DN (Q-T)
        ('Xem', fill_action_dn), ('Thêm', fill_action_dn), ('Sửa', fill_action_dn), ('Xóa', fill_action_dn),
        # Khách (U-W)
        ('Xem', fill_action_khach), ('Thêm/Nộp', fill_action_khach), ('Hủy/Sửa', fill_action_khach),
    ]

    for idx, (act, fill_c) in enumerate(actions, 3): # Col 3 is C
        c = ws.cell(row=5, column=idx, value=act)
        c.fill = fill_c
        c.font = font_header_action
        c.alignment = align_center
        c.border = border_cell

    ws.row_dimensions[5].height = 22

    # Data definition
    # Format: [Type, STT, Chức năng, [Admin: X, T, S, Xoa, Duyet], [So: X, T, S, Xoa, Duyet], [Truong: X, T, S, Xoa], [DN: X, T, S, Xoa], [Khach: X, T, S], Ghi chu]
    # Type: 'G' for Group Header, 'D' for Data Row
    matrix = [
        # Group 1
        ['G', 'I', 'QUẢN LÝ BẢN ĐỒ SỐ GIS & ĐỊA GIỚI HÀNH CHÍNH', None, None, None, None, None, 'Hạ tầng bản đồ Leaflet, OpenFreeMap & ranh giới hành chính'],
        ['D', '1.1', 'Xem bản đồ số, tra cứu vị trí cơ sở GDNN & chỉ đường', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False, False], [True, False, False], 'Công khai cho toàn bộ người dùng tra cứu'],
        ['D', '1.2', 'Quản lý lớp dữ liệu GIS, ranh giới xã/phường/tỉnh Ninh Bình', [True, True, True, True, False], [True, False, True, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Sở & Admin cập nhật dữ liệu sáp nhập địa giới hành chính'],
        ['D', '1.3', 'Phân tích không gian & bán kính phục vụ (Buffer Analysis)', [True, True, True, True, False], [True, True, True, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Phân tích bán kính 3km, 5km, 10km phục vụ quy hoạch'],

        # Group 2
        ['G', 'II', 'QUẢN LÝ CƠ SỞ GIÁO DỤC NGHỀ NGHIỆP (TRƯỜNG HỌC)', None, None, None, None, None, 'Dữ liệu hồ sơ trường, phân hiệu, tọa độ GPS'],
        ['D', '2.1', 'Xem hồ sơ cơ sở GDNN (giới thiệu, ngành học, tọa độ)', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Thông tin công khai trên Web, Mobile App và Zalo'],
        ['D', '2.2', 'Thêm mới cơ sở GDNN vào hệ thống', [True, True, True, True, False], [True, True, True, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Sở GD duyệt mở mới trường hoặc chi nhánh'],
        ['D', '2.3', 'Cập nhật thông tin, hình ảnh, campus cơ sở của trường mình', [True, False, True, False, False], [True, False, True, False, True], [True, False, True, False], [False, False, False, False], [False, False, False], 'Nhà trường chỉ được sửa hồ sơ của chính trường mình'],
        ['D', '2.4', 'Xóa / Đóng mã cơ sở GDNN', [True, False, False, True, False], [True, False, False, True, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Chỉ Sở GD và Admin có quyền ngừng hoạt động cơ sở'],
        ['D', '2.5', 'Import / Export danh sách trường học từ file Excel', [True, True, False, False, False], [True, True, False, False, True], [True, False, False, False], [False, False, False, False], [False, False, False], 'Nhập danh sách hàng loạt bằng file Excel mẫu qua queue nền'],

        # Group 3
        ['G', 'III', 'QUẢN LÝ NGÀNH NGHỀ & TUYỂN SINH', None, None, None, None, None, 'Danh mục nghề đào tạo, chỉ tiêu và hồ sơ tư vấn'],
        ['D', '3.1', 'Tra cứu danh mục ngành nghề, chương trình đào tạo & học phí', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Công khai cho học sinh và doanh nghiệp tìm kiếm'],
        ['D', '3.2', 'Quản lý danh mục ngành nghề chuẩn (8 nhóm ngành GDNN)', [True, True, True, True, False], [True, True, True, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Sở GD chuẩn hóa danh mục nghề cấp tỉnh'],
        ['D', '3.3', 'Cập nhật chỉ tiêu tuyển sinh, học phí, mô tả nghề của trường', [True, False, True, False, False], [True, False, True, False, True], [True, True, True, True], [False, False, False, False], [False, False, False], 'Nhà trường tự quản lý chỉ tiêu & thông báo tuyển sinh của mình'],
        ['D', '3.4', 'Nộp hồ sơ / Đăng ký tư vấn tuyển sinh trực tuyến', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [False, False, False, False], [False, True, True], 'Học sinh/người dân gửi form nguyện vọng trực tiếp cho trường'],
        ['D', '3.5', 'Tiếp nhận, xử lý & phản hồi hồ sơ đăng ký tư vấn tuyển sinh', [True, False, False, False, False], [True, False, False, False, False], [True, False, True, False], [False, False, False, False], [False, False, False], 'Cán bộ tuyển sinh của trường liên hệ và xử lý'],

        # Group 4
        ['G', 'IV', 'QUẢN LÝ VIỆC LÀM & TUYỂN DỤNG LAO ĐỘNG', None, None, None, None, None, 'Cầu nối Doanh nghiệp - Học sinh, sinh viên - Người lao động'],
        ['D', '4.1', 'Xem danh sách tin tuyển dụng, lọc theo KCN & mức lương', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Hiển thị công khai trên Web, App và Zalo Mini App'],
        ['D', '4.2', 'Đăng tin tuyển dụng mới', [True, True, True, True, False], [True, False, False, False, True], [False, False, False, False], [True, True, False, False], [False, False, False], 'Doanh nghiệp đăng tin tuyển dụng vị trí việc làm'],
        ['D', '4.3', 'Chỉnh sửa, cập nhật, gia hạn tin tuyển dụng', [True, False, True, False, False], [True, False, False, False, True], [False, False, False, False], [True, False, True, False], [False, False, False], 'Doanh nghiệp quản lý các tin đăng của chính mình'],
        ['D', '4.4', 'Xóa / Đóng tin tuyển dụng hết hạn', [True, False, False, True, False], [True, False, False, True, True], [False, False, False, False], [True, False, False, True], [False, False, False], 'Doanh nghiệp tự đóng; Sở/Admin hạ tin vi phạm'],
        ['D', '4.5', 'Kiểm duyệt / Phê duyệt tin tuyển dụng của Doanh nghiệp', [True, False, False, False, True], [True, False, False, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Kiểm duyệt nội dung tin đăng tránh lừa đảo, đa cấp'],
        ['D', '4.6', 'Nộp hồ sơ ứng tuyển việc làm trực tuyến (Gửi CV)', [True, False, False, False, False], [False, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, True, True], 'Học viên/người tìm việc nộp hồ sơ trực tuyến'],
        ['D', '4.7', 'Quản lý danh sách ứng viên & xem hồ sơ ứng tuyển', [True, False, False, False, False], [False, False, False, False, False], [False, False, False, False], [True, False, True, False], [False, False, False], 'Doanh nghiệp quản lý hồ sơ ứng viên gửi về cho các bài đăng'],

        # Group 5
        ['G', 'V', 'QUẢN LÝ DOANH NGHIỆP & LIÊN KẾT ĐÀO TẠO', None, None, None, None, None, 'Hồ sơ doanh nghiệp, KCN và kết nối cung - cầu lao động'],
        ['D', '5.1', 'Xem hồ sơ thông tin doanh nghiệp, nhà máy trong KCN', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Công khai thông tin doanh nghiệp đối tác'],
        ['D', '5.2', 'Đăng ký & Cập nhật hồ sơ thông tin doanh nghiệp', [True, False, True, False, False], [True, False, True, False, True], [False, False, False, False], [True, True, True, False], [False, False, False], 'Doanh nghiệp tự đăng ký tài khoản và cập nhật hồ sơ'],
        ['D', '5.3', 'Phê duyệt tài khoản & hồ sơ doanh nghiệp', [True, False, False, False, True], [True, False, False, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Sở GD/Sở LĐ-TB&XH xác minh pháp nhân doanh nghiệp'],
        ['D', '5.4', 'Quản lý danh mục KCN & Cụm công nghiệp trên địa bàn', [True, True, True, True, False], [True, True, True, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Dữ liệu các KCN (Gián Khẩu, Phúc Sơn, Tam Điệp, Khánh Phú...)'],
        ['D', '5.5', 'Gửi yêu cầu liên kết đào tạo giữa Nhà trường & Doanh nghiệp', [True, False, False, False, False], [True, False, False, False, False], [True, True, True, False], [True, True, True, False], [False, False, False], 'Ký kết tiếp nhận thực tập sinh, đặt hàng đào tạo'],

        # Group 6
        ['G', 'VI', 'BÁO CÁO THỐNG KÊ & PHÂN TÍCH SỐ LIỆU', None, None, None, None, None, 'Báo cáo tổng quan, số liệu tuyển sinh và thị trường lao động'],
        ['D', '6.1', 'Xem Dashboard thống kê tổng quan toàn tỉnh', [True, False, False, False, False], [True, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Tổng số trường, ngành nghề, việc làm, mật độ phân bố'],
        ['D', '6.2', 'Thống kê cơ sở GDNN theo địa bàn huyện/thành phố/xã', [True, False, False, False, False], [True, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Phục vụ công tác quy hoạch mạng lưới GDNN của tỉnh'],
        ['D', '6.3', 'Báo cáo tuyển sinh & lượt quan tâm của trường mình', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [False, False, False, False], [False, False, False], 'Nhà trường theo dõi hiệu quả tuyển sinh'],
        ['D', '6.4', 'Báo cáo lượt xem tin tuyển dụng & ứng viên của doanh nghiệp', [True, False, False, False, False], [True, False, False, False, False], [False, False, False, False], [True, False, False, False], [False, False, False], 'Doanh nghiệp theo dõi hiệu quả thu hút nhân lực'],
        ['D', '6.5', 'Xuất báo cáo dữ liệu định kỳ (Excel / PDF)', [True, True, False, False, False], [True, True, False, False, False], [True, True, False, False], [True, True, False, False], [False, False, False], 'Xuất các biểu mẫu báo cáo thống kê phục vụ hội nghị'],

        # Group 7
        ['G', 'VII', 'TƯƠNG TÁC, PHẢN ÁNH & TIN TỨC', None, None, None, None, None, 'Kênh tiếp nhận ý kiến nhân dân và bản tin tuyển sinh/việc làm'],
        ['D', '7.1', 'Xem bài viết tin tức, hướng nghiệp, thông báo của ngành', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Công khai toàn dân tra cứu'],
        ['D', '7.2', 'Đăng bài viết tin tức, thông báo tuyển sinh', [True, True, True, True, False], [True, True, True, True, True], [True, True, True, False], [False, False, False, False], [False, False, False], 'Nhà trường đăng tin của trường; Sở đăng tin toàn tỉnh'],
        ['D', '7.3', 'Gửi phản ánh, góp ý về cơ sở đào tạo hoặc tin tuyển dụng', [True, False, False, False, False], [True, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, True, False], 'Người dân gửi phản ánh chất lượng đào tạo / tuyển dụng sai'],
        ['D', '7.4', 'Tiếp nhận, xử lý & phản hồi phản ánh của người dân', [True, False, True, False, False], [True, False, True, False, True], [False, False, False, False], [False, False, False, False], [False, False, False], 'Sở GD chỉ đạo xác minh và trả lời người dân'],

        # Group 8
        ['G', 'VIII', 'QUẢN TRỊ HỆ THỐNG & TÀI KHOẢN', None, None, None, None, None, 'Bảo mật, phân quyền, cấu hình hạ tầng và nhật ký hệ thống'],
        ['D', '8.1', 'Đăng nhập hệ thống (Web / Mobile App / Zalo Mini App)', [True, False, False, False, False], [True, False, False, False, False], [True, False, False, False], [True, False, False, False], [True, False, False], 'Khách có thể đăng nhập bằng Google/Zalo hoặc xem ẩn danh'],
        ['D', '8.2', 'Quản trị tài khoản người dùng & phân quyền 5 vai trò', [True, True, True, True, False], [True, False, True, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Admin quản lý toàn bộ tài khoản; Sở xem danh sách cán bộ'],
        ['D', '8.3', 'Cấu hình hệ thống (Banner, giao diện, API, cấu hình email/SMS)', [True, True, True, True, False], [False, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Dành riêng cho Quản trị viên kỹ thuật'],
        ['D', '8.4', 'Xem nhật ký hệ thống (Audit Log & Activity Tracking)', [True, False, False, False, False], [True, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Ghi lại mọi thao tác thêm/sửa/xóa để phục vụ kiểm toán'],
        ['D', '8.5', 'Sao lưu & phục hồi CSDL (Backup & Restore)', [True, True, False, True, False], [False, False, False, False, False], [False, False, False, False], [False, False, False, False], [False, False, False], 'Tự động sao lưu database hàng ngày lên Cloudflare R2/S3']
    ]

    current_row = 6
    for item in matrix:
        ws.row_dimensions[current_row].height = 20
        if item[0] == 'G':
            # Group Header
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=2)
            c_lbl = ws.cell(row=current_row, column=1, value=f"{item[1]}. {item[2]}")
            c_lbl.font = font_group
            c_lbl.alignment = align_left

            for col in range(1, 25):
                cell = ws.cell(row=current_row, column=col)
                cell.fill = fill_group_row
                cell.border = border_cell

            # Note column
            c_note = ws.cell(row=current_row, column=24, value=item[8])
            c_note.font = font_note
            c_note.alignment = align_left
        else:
            # Data row
            stt = item[1]
            func_name = item[2]
            admin_acts = item[3]
            so_acts = item[4]
            truong_acts = item[5]
            dn_acts = item[6]
            khach_acts = item[7]
            note = item[8]

            # STT
            c_stt = ws.cell(row=current_row, column=1, value=stt)
            c_stt.alignment = align_center
            c_stt.font = font_text
            c_stt.border = border_cell

            # Name
            c_name = ws.cell(row=current_row, column=2, value=func_name)
            c_name.alignment = align_left
            c_name.font = font_text
            c_name.border = border_cell

            # All bools
            all_bools = admin_acts + so_acts + truong_acts + dn_acts + khach_acts
            for b_idx, val in enumerate(all_bools, 3):
                c_val = ws.cell(row=current_row, column=b_idx)
                c_val.value = "TRUE" if val else "FALSE"
                c_val.alignment = align_center
                c_val.border = border_cell
                if val:
                    c_val.font = font_true
                    c_val.fill = fill_true
                else:
                    c_val.font = font_false
                    if current_row % 2 == 0:
                        c_val.fill = fill_zebra

            # Note
            c_n = ws.cell(row=current_row, column=24, value=note)
            c_n.font = font_note
            c_n.alignment = align_left
            c_n.border = border_cell

        current_row += 1

    # Column widths
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 46

    # Action columns C to W
    for c_idx in range(3, 24):
        ws.column_dimensions[get_column_letter(c_idx)].width = 7.5

    ws.column_dimensions['X'].width = 45
    ws.views.sheetView[0].showGridLines = True

    wb.save(output_path)
    print(f"Successfully generated permission matrix at: {output_path}")

if __name__ == '__main__':
    create_permission_matrix_excel('Bang-phan-quyen-du-an-SchoolMap.xlsx')
